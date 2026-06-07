"""
FAQ (Frequently Asked Questions) extraction service.
Extracts Q&A pairs from ASR transcripts using LLM analysis.
"""

import os
import json
import re
from typing import List, Dict, Optional
from models.schemas import ASRResult, LLMConfig
from services.llm_service import call_llm, call_llm_with_json_output, extract_json_from_text
from services.prompt_service import get_prompt


def extract_faqs(
    asr_results: List[ASRResult],
    llm_config: LLMConfig,
) -> List[Dict]:
    """Extract FAQ pairs from ASR results."""
    # Build combined transcript
    parts = []
    for result in asr_results:
        audio_name = os.path.splitext(result.filename)[0]
        parts.append(f"\n=== 音频文件: {audio_name} ===\n")
        if result.sentences:
            for sent in result.sentences:
                speaker_prefix = f"[{sent.speaker}] " if sent.speaker else ""
                parts.append(f"第{sent.index}句话：{speaker_prefix}{sent.text}")
        else:
            parts.append(f"（全文）{result.full_text}")
        parts.append("")

    combined_transcript = "\n".join(parts)

    # Use prompt from prompt_service (user-customizable)
    prompt = get_prompt("faq_extraction")
    if not prompt or '{transcripts}' not in prompt:
        prompt = FAQ_EXTRACTION_PROMPT
    prompt_filled = prompt.replace('{transcripts}', combined_transcript)

    result = call_llm_with_json_output(
        messages=[{"role": "user", "content": prompt_filled}],
        config=llm_config
    )

    data = extract_json_from_text(result) or {}
    faqs = data.get("faqs", [])

    # Deduplicate FAQ entries
    seen = set()
    unique_faqs = []
    for faq in faqs:
        q = faq.get("question", "").strip()
        # Normalize for dedup
        q_norm = re.sub(r'\s+', '', q.lower())
        if q_norm and q_norm not in seen:
            seen.add(q_norm)
            unique_faqs.append(faq)

    return unique_faqs


FAQ_EXTRACTION_PROMPT = """你是一个专业的FAQ提取专家。请从以下电话录音的ASR转录文本中提取FAQ（常见问题与解答）。

## 转录文本
{transcripts}

## 提取要求
请提取对话中出现的所有"问题-回答"对，包括客户问题和客服解答。

## 去重规则
相同或高度相似的问题只保留一个，保留最完整的解答。

## 输出JSON格式
{"faqs": [{"question": "...", "answer": "...", "category": "...", "source_file": "..."}]}

请确保不重复提取相同的FAQ。输出完整JSON。"""


def generate_faq_excel(faqs: List[Dict], output_path: str) -> str:
    """Generate an Excel file from FAQ data."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "FAQ"

        # Header styling
        header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF', size=12)
        header_fill = PatternFill(start_color='00a986', end_color='00a986', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # Headers
        headers = ['序号', '问题分类', '客户问题', '客服解答', '来源音频']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Data
        for i, faq in enumerate(faqs, 1):
            row_data = [
                i,
                faq.get('category', ''),
                faq.get('question', ''),
                faq.get('answer', ''),
                faq.get('source_file', '')
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=i + 1, column=col, value=value)
                cell.font = Font(name='Microsoft YaHei', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = thin_border

        # Column widths
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 40
        ws.column_dimensions['D'].width = 60
        ws.column_dimensions['E'].width = 30

        wb.save(output_path)
        return output_path
    except ImportError:
        # Fallback: generate CSV
        import csv
        csv_path = output_path.replace('.xlsx', '.csv')
        with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['序号', '问题分类', '客户问题', '客服解答', '来源音频'])
            for i, faq in enumerate(faqs, 1):
                writer.writerow([
                    i,
                    faq.get('category', ''),
                    faq.get('question', ''),
                    faq.get('answer', ''),
                    faq.get('source_file', '')
                ])
        return csv_path